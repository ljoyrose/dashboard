"""Slide engine. A deck is a JSON file (course/decks/<deck>.json): a list of slides, each with a layout and its words.
One layout engine turns each slide into drawing parts; two writers turn the parts into
  1. SVG per slide (-> PNG and one vector PDF via shoot.mjs), and
  2. an editable .pptx (native text boxes and shapes; icons as small pictures).
So the PowerPoint and the PDF always match. Brand jackets (LJR7, PPP, Your Brand) swap tokens only (D8).

Usage: python build/deck.py decks/keynote.json [--brand ppp]
"""
import json
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(__file__))
import lj
from lj import G, C, tw, wrap, esc

W, H = 1920, 1080
GUT = G["space"]["gutter_1920"]
FOOT_H = 84
T = G["type"]["slide_1920"]
LH = G["type"]["line_height"]


# ---------- drawing parts ----------
class Parts(list):
    def rect(self, x, y, w, h, fill, stroke=None, r=0, shadow=False):
        self.append(("rect", x, y, w, h, fill, stroke, r, shadow))

    def card(self, x, y, w, h, accent=None):
        r = G["space"]["radius"]
        self.rect(x, y, w, h, C["card"], C["border"], r, True)
        if accent:
            self.append(("band", x, y, w, h, accent, G["space"]["band"], r))

    def text(self, x, y, w, rows, size, weight=400, colour=None, align="l"):
        if isinstance(rows, str):
            rows = wrap(rows, size, w, weight)
        self.append(("text", x, y, w, rows, size, weight, colour or C["ink"], align))
        return len(rows) * size * LH

    def icon(self, name, x, y, size, colour=None):
        self.append(("icon", name, x, y, size, colour or C["ink"]))

    def image(self, path, x, y, w, h):
        self.append(("image", path, x, y, w, h))

    def pill(self, x, y, label, size=26, icon_name=None, colour=None):
        h = int(size * 1.7)
        pad = int(size * 0.6)
        iw = int(size * 1.05) + 6 if icon_name else 0
        w = int(tw(label, size, 600) + 2 * pad + iw)
        self.rect(x, y, w, h, C["card"], C["border"], h / 2)
        if icon_name:
            self.icon(icon_name, x + pad, y + (h - size) / 2, size, colour)
        self.text(x + pad + iw, y + (h - size * LH) / 2, w, [label], size, 600, colour)
        return w


# ---------- frame ----------
def frame(p, s, meta):
    top = 64
    p.text(GUT, top, W - 2 * GUT - 420, [s["title"]], T["title"] if len(s["title"]) < 34 else 64, 700)
    th = (T["title"] if len(s["title"]) < 34 else 64) * LH
    if s.get("sub"):
        p.text(GUT, top + th + 2, W - 2 * GUT, [s["sub"]], 32, 400, C["muted"])
    rule_y = top + th + 2 + (32 * LH if s.get("sub") else 0) + 16
    p.rect(GUT, rule_y, W - 2 * GUT, 2, C["ink"])
    if s.get("vision") or meta.get("vision"):
        label = G["brand"]["vision_label"]
        wv = int(tw(label, 26, 600) + 26 * 1.2 + 26 * 1.05 + 6)
        p.pill(W - GUT - wv, top + 8, label, 26, "sparkles")
    elif s.get("chip"):
        wv = int(tw(s["chip"], 26, 600) + 26 * 1.2)
        p.pill(W - GUT - wv, top + 8, s["chip"], 26)
    # footer band (rule 3)
    p.rect(0, H - FOOT_H, W, FOOT_H, C["ink"])
    date = G["brand"]["course_date"]
    p.text(GUT, H - FOOT_H + 12, 400, [date], 26, 700, C["footer_text"])
    p.text(GUT + tw(date, 26, 700) + 28, H - FOOT_H + 13, 900, [meta["module"] + f" · Slide {s['_n']}"], 24, 400, C["footer_text"])
    brand = meta["brand"]
    p.text(GUT, H - FOOT_H + 48, W - 2 * GUT - 400, [s.get("source", meta.get("source", ""))], 22, 400, C["border"])
    p.text(W - GUT - 500, H - FOOT_H + 22, 500, [brand["long_name"]], 24, 600, C["footer_text"], "r")
    return rule_y + 2


def key_row(p, items, y):
    """Key above the footer (rule 10). items: [(icon, label)]. Returns top y used."""
    if not items:
        return y
    ky = H - FOOT_H - 64
    p.rect(GUT, ky - 14, W - 2 * GUT, 1, C["border"])
    x = GUT
    p.text(x, ky, 120, ["Key"], 24, 700)
    x += 90
    for ic, label in items:
        if ic.startswith("#"):
            p.rect(x, ky + 6, 26, 22, ic, None, 4)
        elif ic.startswith("status:"):
            st = next(t for t in G["statuses"] if t["key"] == ic[7:])
            p.icon(st["icon"], x, ky + 2, 28, C["status"][st["key"]])
        else:
            p.icon(ic, x, ky + 2, 28)
        x += 40
        p.text(x, ky + 2, 600, [label], 24, 400, C["muted"])
        x += tw(label, 24) + 44
    return ky - 20


# ---------- layouts ----------
def L_title(p, s, top, bottom):
    y = top + 120
    y += p.text(GUT, y, 1100, s.get("big", s["title"]), 104, 700)
    if s.get("line"):
        y += 20 + p.text(GUT, y + 20, 1000, s["line"], 40, 400, C["muted"])
    if s.get("image"):
        p.image(s["image"], 1180, top + 40, 640, bottom - top - 80)
    elif s.get("icon"):
        p.icon(s["icon"], 1350, top + 160, 360, s.get("icon_colour"))


def L_statement(p, s, top, bottom):
    if s.get("icon"):
        p.icon(s["icon"], GUT, top + 110, 160, s.get("icon_colour"))
    x = GUT + (230 if s.get("icon") else 0)
    y = top + 100
    y += p.text(x, y, W - x - GUT, s["statement"], T["statement"], 700)
    if s.get("line"):
        p.text(x, y + 30, W - x - GUT, s["line"], 40, 400, C["muted"])


def L_cards(p, s, top, bottom):
    cards = s["cards"]
    n = len(cards)
    cols = n if n <= 4 else (3 if n in (5, 6, 9) else 4)
    rows = (n + cols - 1) // cols
    gap = 32
    cw = (W - 2 * GUT - gap * (cols - 1)) / cols
    avail = bottom - top - 60
    ch = min(380, (avail - gap * (rows - 1)) / rows)
    y0 = top + 40
    for i, c in enumerate(cards):
        r, k = divmod(i, cols)
        x, y = GUT + k * (cw + gap), y0 + r * (ch + gap)
        acc = lj.layer(c["layer"]) if c.get("layer") else (lj.ring(c["ring"]) if c.get("ring") else None)
        p.card(x, y, cw, ch, acc)
        iy = y + 36
        if c.get("icon"):
            p.icon(c["icon"], x + 32, iy, 56 if ch > 250 else 44)
            iy += (56 if ch > 250 else 44) + 22
        tsz = 40 if cw > 500 else 34
        iy += p.text(x + 32, iy, cw - 64, c["title"], tsz, 700) + 6
        if c.get("line"):
            p.text(x + 32, iy, cw - 64, c["line"], 28 if cw > 380 else 26, 400, C["muted"])


def L_steps(p, s, top, bottom):
    steps = s["steps"]
    n = len(steps)
    gap = 18
    cw = (W - 2 * GUT - gap * (n - 1)) / n
    y = top + 70
    ch = 300 if n <= 6 else 280
    for i, st in enumerate(steps):
        x = GUT + i * (cw + gap)
        acc = lj.layer(st["layer"]) if st.get("layer") else None
        p.card(x, y, cw, ch, acc)
        isz = 64 if cw > 200 else 48
        p.icon(st["icon"], x + (cw - isz) / 2, y + 36, isz)
        if st.get("n") is not None:
            p.text(x + 14, y + 16, 60, [str(st["n"])], 28, 700, C["muted"])
        tsz = 30 if cw > 200 else 24
        p.text(x + 12, y + 36 + isz + 24, cw - 24, wrap(st["label"], tsz, cw - 24, 700), tsz, 700, None, "c")
        if st.get("gate"):
            g = st["gate"]
            gy = y + ch + 24
            human = g.get("human")
            p.icon("user" if human else "robot", x + 14, gy + 4, 28, C["ink"] if human else C["muted"])
            p.text(x + 50, gy + 2, cw - 50, [g["name"]], 26, 600, C["ink"] if human else C["muted"])
    if s.get("line"):
        p.text(GUT, y + ch + (110 if any(st.get("gate") for st in steps) else 60), W - 2 * GUT, s["line"], 34, 400, C["muted"])


def L_image(p, s, top, bottom):
    import PIL.Image
    im = PIL.Image.open(s["image"])
    iw, ih = im.size
    box_w = W - 2 * GUT - (620 if s.get("points") else 0)
    box_h = bottom - top - 50
    sc = min(box_w / iw, box_h / ih)
    w, h = iw * sc, ih * sc
    p.rect(GUT - 1, top + 30 - 1, w + 2, h + 2, C["card"], C["border"], 8, True)
    p.image(s["image"], GUT, top + 30, w, h)
    if s.get("points"):
        x = GUT + w + 60
        y = top + 60
        for pt in s["points"]:
            p.icon(pt.get("icon", "arrow-right"), x, y + 4, 36)
            y += p.text(x + 56, y, W - GUT - x - 56, pt["text"], 32, 600) + 26


def L_recap(p, s, top, bottom):
    y = top + 50
    for pt in s["points"]:
        p.icon("circle-check", GUT, y, 48, C["status"]["solid"])
        y += max(60, p.text(GUT + 76, y + 4, W - 2 * GUT - 76, pt, 38, 600)) + 20
    if s.get("why"):
        y += 20
        rows = wrap(s["why"], 34, W - 2 * GUT - 80, 700)
        hgt = 70 + len(rows) * 34 * LH + 20
        p.rect(GUT, y, W - 2 * GUT, hgt, C["sunk"], None, G["space"]["radius"])
        p.rect(GUT, y, 6, hgt, C["ink"], None, 3)
        p.text(GUT + 40, y + 24, 600, ["WHY IT MATTERS"], 22, 700, C["muted"])
        p.text(GUT + 40, y + 62, W - 2 * GUT - 80, rows, 34, 700)


def L_quiz(p, s, top, bottom):
    y = top + 40
    y += p.text(GUT, y, W - 2 * GUT, s["q"], 48, 700) + 30
    opts = s["options"]
    cw = (W - 2 * GUT - 32) / 2
    ch = 150
    for i, o in enumerate(opts):
        r, k = divmod(i, 2)
        x, yy = GUT + k * (cw + 32), y + r * (ch + 28)
        p.card(x, yy, cw, ch)
        p.text(x + 32, yy + 44, 60, ["ABCD"[i]], 44, 700, C["muted"])
        p.text(x + 100, yy + 30, cw - 130, o, 32, 600)


def L_cta(p, s, top, bottom):
    y = top + 50
    for i, st in enumerate(s["steps"], 1):
        p.rect(GUT, y, 64, 64, C["card"], C["border"], 32)
        p.text(GUT, y + 10, 64, [str(i)], 32, 700, None, "c")
        y += max(80, p.text(GUT + 96, y + 8, 1100, st, 38, 600) + 24)
    x = 1340
    yy = top + 60
    for lk in s.get("links", []):
        p.card(x, yy, W - GUT - x, 150)
        p.icon(lk.get("icon", "link"), x + 30, yy + 30, 44)
        p.text(x + 96, yy + 28, W - GUT - x - 120, [lk["label"]], 30, 700)
        p.text(x + 96, yy + 76, W - GUT - x - 120, lk["value"], 24, 400, C["muted"])
        yy += 180
    if s.get("note"):
        p.text(GUT, bottom - 70, W - 2 * GUT, s["note"], 26, 400, C["muted"])


def L_table(p, s, top, bottom):
    cols = s["columns"]
    rows = s["rows"]
    n = len(cols)
    widths = s.get("widths") or [1] * n
    tot = sum(widths)
    xs = [GUT]
    for w_ in widths:
        xs.append(xs[-1] + (W - 2 * GUT) * w_ / tot)
    y = top + 40
    rh = min(76, (bottom - top - 140) / (len(rows) + 1))
    p.card(GUT, y, W - 2 * GUT, rh * (len(rows) + 1) + 10)
    for i, c in enumerate(cols):
        p.text(xs[i] + 24, y + (rh - 28 * LH) / 2 + 4, xs[i + 1] - xs[i] - 30, [c], 28, 700)
    for r_, row in enumerate(rows, 1):
        yy = y + r_ * rh
        p.rect(GUT + 16, yy, W - 2 * GUT - 32, 1, C["border"])
        for i, cell in enumerate(row):
            if isinstance(cell, dict) and cell.get("status"):
                st = next(t for t in G["statuses"] if t["key"] == cell["status"])
                p.icon(st["icon"], xs[i] + 24, yy + (rh - 30) / 2, 30, C["status"][st["key"]])
                p.text(xs[i] + 62, yy + (rh - 28 * LH) / 2 + 2, xs[i + 1] - xs[i] - 70, [st["word"]], 28, 600, C["status"][st["key"]])
            else:
                p.text(xs[i] + 24, yy + (rh - 28 * LH) / 2 + 2, xs[i + 1] - xs[i] - 30, [str(cell)], 28, 600 if i == 0 else 400)


LAYOUTS = {"title": L_title, "statement": L_statement, "cards": L_cards, "steps": L_steps, "image": L_image,
           "recap": L_recap, "quiz": L_quiz, "cta": L_cta, "table": L_table}


def build_parts(s, meta):
    p = Parts()
    top = frame(p, s, meta)
    bottom = key_row(p, s.get("key", []), H - FOOT_H - 24)
    LAYOUTS[s["layout"]](p, s, top, bottom)
    return p


# ---------- writers ----------
def to_svg(parts):
    body = []
    for pt in parts:
        k = pt[0]
        if k == "rect":
            _, x, y, w, h, fill, stroke, r, shadow = pt
            body.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="{r}" fill="{fill}"'
                        + (f' stroke="{stroke}" stroke-width="1"' if stroke else "")
                        + (' filter="url(#soft)"' if shadow else "") + "/>")
        elif k == "text":
            _, x, y, w, rows, size, weight, colour, align = pt
            anchor, tx = {"l": ("start", x), "c": ("middle", x + w / 2), "r": ("end", x + w)}[align]
            for i, row in enumerate(rows):
                by = y + i * size * LH + size * 0.93
                body.append(lj.text(tx, by, row, size, weight, colour, anchor))
        elif k == "band":
            _, x, y, w, h, colour, band, r = pt
            body.append(lj.band_svg(x, y, w, h, colour, band, r))
        elif k == "icon":
            _, name, x, y, size, colour = pt
            body.append(lj.icon(name, x, y, size, colour))
        elif k == "image":
            _, path, x, y, w, h = pt
            body.append(f'<image href="{esc(os.path.abspath(path))}" x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" preserveAspectRatio="xMidYMid meet"/>')
    return lj.svg_doc(W, H, "".join(body))


def icon_pngs(decks_parts, cache_dir):
    """Render every (icon, colour) used, once, as a 192 px PNG for the pptx."""
    need = set()
    for parts in decks_parts:
        for pt in parts:
            if pt[0] == "icon":
                need.add((pt[1], pt[5]))
    todo = []
    for name, colour in need:
        f = os.path.join(cache_dir, f"{name}-{colour[1:]}.svg")
        if not os.path.exists(f.replace(".svg", ".png")):
            os.makedirs(cache_dir, exist_ok=True)
            open(f, "w").write(f'<svg xmlns="http://www.w3.org/2000/svg" width="192" height="192" viewBox="0 0 192 192">{lj.icon(name, 0, 0, 192, colour)}</svg>')
            todo.append(f)
    if todo:
        subprocess.run(["node", os.path.join(lj.HERE, "shoot.mjs"), "--transparent"] + todo, check=True, stdout=subprocess.DEVNULL)


def to_pptx(slides_parts, notes, out, cache_dir):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
    from lxml import etree

    E = 6350  # EMU per px on a 1920 x 1080 canvas
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W * E), Emu(H * E)
    blank = prs.slide_layouts[6]
    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    for parts, note in zip(slides_parts, notes):
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(C["page"])
        for pt in parts:
            k = pt[0]
            if k == "rect":
                _, x, y, w, h, fill, stroke, r, shadow = pt
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,
                                          Emu(int(x * E)), Emu(int(y * E)), Emu(int(w * E)), Emu(int(h * E)))
                if r:
                    shp.adjustments[0] = min(0.5, r / max(1, min(w, h)))
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(fill)
                if stroke:
                    shp.line.color.rgb = rgb(stroke)
                    shp.line.width = Emu(E)
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
                shp.text_frame.text = ""
            elif k == "text":
                _, x, y, w, rows, size, weight, colour, align = pt
                tb = sl.shapes.add_textbox(Emu(int(x * E)), Emu(int(y * E)), Emu(int((w + 8) * E)), Emu(int(len(rows) * size * LH * E + 4 * E)))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, row in enumerate(rows):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
                    para.line_spacing = Pt(size * LH * 0.5)
                    run = para.add_run()
                    run.text = row
                    f = run.font
                    f.name = "Lexend" if weight < 600 else ("Lexend SemiBold" if weight == 600 else "Lexend")
                    f.size = Pt(size * 0.5)
                    f.bold = weight >= 700
                    f.color.rgb = rgb(colour)
            elif k == "band":
                _, x, y, w, h, colour, band, r = pt
                ins = r * 0.45
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, Emu(int((x + ins) * E)), Emu(int(y * E)),
                                          Emu(int((w - 2 * ins) * E)), Emu(int(band * E)))
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(colour)
                shp.line.fill.background()
                shp.shadow.inherit = False
            elif k == "icon":
                _, name, x, y, size, colour = pt
                png = os.path.join(cache_dir, f"{name}-{colour[1:]}.png")
                sl.shapes.add_picture(png, Emu(int(x * E)), Emu(int(y * E)), Emu(int(size * E)), Emu(int(size * E)))
            elif k == "image":
                _, path, x, y, w, h = pt
                sl.shapes.add_picture(path, Emu(int(x * E)), Emu(int(y * E)), Emu(int(w * E)), Emu(int(h * E)))
        if note:
            sl.notes_slide.notes_text_frame.text = note
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)


def jacket(brand_key):
    return G["brands"][brand_key]


def render_deck(spec_path, brand_key=None, out_dir=None, pdf=True):
    spec = json.load(open(spec_path, encoding="utf-8"))
    brand_key = brand_key or spec.get("brand", "ljr7")
    meta = {"module": spec["module"], "brand": jacket(brand_key), "source": spec.get("source", ""),
            "vision": spec.get("vision", False)}
    name = spec["name"] + ("" if brand_key == "ljr7" else f"-{brand_key.upper()}")
    out_dir = out_dir or os.path.join(lj.ROOT, "slides", spec["name"])
    svg_dir = os.path.join(out_dir, "svg" if brand_key == "ljr7" else f"svg-{brand_key}")
    os.makedirs(svg_dir, exist_ok=True)
    all_parts, notes, svgs = [], [], []
    for i, s in enumerate(spec["slides"], 1):
        s = dict(s)
        s["_n"] = i
        s = brandify(s, meta["brand"])
        parts = build_parts(s, meta)
        all_parts.append(parts)
        notes.append(s.get("notes", "") + (f"\n\nFramework beat: {s['beat']}" if s.get("beat") else ""))
        f = os.path.join(svg_dir, f"{name}-{i:02d}.svg")
        lj.save(f, to_svg(parts))
        svgs.append(f)
    cache = os.path.join(lj.ROOT, "build", ".iconcache")
    icon_pngs(all_parts, cache)
    pptx = os.path.join(out_dir, f"{name}-{G['brand']['module_version']}.pptx")
    to_pptx(all_parts, notes, pptx, cache)
    if pdf:
        subprocess.run(["node", os.path.join(lj.HERE, "shoot.mjs"), "--pdf", os.path.join(out_dir, f"{name}-{G['brand']['module_version']}.pdf")] + svgs, check=True, stdout=subprocess.DEVNULL)
        subprocess.run(["node", os.path.join(lj.HERE, "shoot.mjs")] + svgs, check=True, stdout=subprocess.DEVNULL)
    return pptx, svgs


def brandify(s, brand):
    """Replace {brand}, {brand_long}, {join} in every string of a slide."""
    def sub(o):
        if isinstance(o, str):
            return (o.replace("{brand}", brand["name"]).replace("{brand_long}", brand["long_name"])
                    .replace("{join}", brand["join_link"]))
        if isinstance(o, list):
            return [sub(x) for x in o]
        if isinstance(o, dict):
            return {k: sub(v) for k, v in o.items()}
        return o
    return sub(s)


if __name__ == "__main__":
    args = sys.argv[1:]
    brand = None
    if "--brand" in args:
        brand = args[args.index("--brand") + 1]
        args = [a for a in args if a not in ("--brand", brand)]
    for a in args:
        print(render_deck(a, brand)[0])
